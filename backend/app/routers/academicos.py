import io
import re
import unicodedata
import openpyxl
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, status, Query, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Optional
from app.db.session import SessionLocal
from app.models.models import (
    Estudiante, Grupo, Materia, Periodo, Horario, Docente, Carrera,
    Usuario, DirectorCarrera, DiaSemanaEnum, HistorialAcademicoPrevio, Tutor, RolEnum
)
from app.core.deps import get_db, RoleChecker, get_current_active_user
from app.schemas.academic import (
    EstudianteCreate, EstudianteOut, GrupoCreate, GrupoOut, 
    MateriaCreate, MateriaOut, MateriaUpdate, PeriodoCreate, PeriodoOut,
    ClaseDocenteOut, EstudianteBajaIn,
    HistorialAcademicoPrevioCreate, HistorialAcademicoPrevioOut,
    EstudianteAltaCreate
)
from app.core.deps import get_db, RoleChecker, get_current_active_user
from app.schemas.carrera import CarreraOut
from fastapi.responses import StreamingResponse
from app.schemas.horario import HorarioBulkCreate, HorarioOut, DocenteCatalogoOut

router = APIRouter(prefix="/api/v1", tags=["Académico"])

solo_admin = RoleChecker(["Administrador", "Psicopedagogia"])
permitir_gestion_materias = RoleChecker(["Administrador", "Director", "Psicopedagogia"])
todos_los_roles = RoleChecker(["Administrador", "Tutor", "Docente", "Director", "Psicopedagogia", "RRHH", "Mixto"])
permitir_gestion_grupos = RoleChecker(["Administrador", "Director", "Psicopedagogia"])

def _generar_matricula(db: Session) -> str:
    anio = datetime.now().year
    prefijo = f"MAT-{anio}-"
    ultimo = db.query(Estudiante)\
        .filter(Estudiante.matricula.like(f"{prefijo}%"))\
        .order_by(Estudiante.matricula.desc())\
        .first()
    if ultimo:
        try:
            consecutivo = int(ultimo.matricula.split('-')[-1]) + 1
        except (ValueError, IndexError):
            consecutivo = 1
    else:
        consecutivo = 1
    return f"{prefijo}{consecutivo:04d}"

def _armar_materia_salida(m: Materia, db: Session) -> MateriaOut:
    carrera = db.query(Carrera).filter(Carrera.id_carrera == m.id_carrera).first() if m.id_carrera else None
    return MateriaOut(
        id_materia=m.id_materia,
        nombre_materia=m.nombre_materia,
        clave_materia=m.clave_materia,
        creditos=m.creditos,
        horas_semana=m.horas_semana,
        id_carrera=m.id_carrera,
        nombre_carrera=carrera.nombre if carrera else None,
        estado=m.estado
    )

def _generar_correo_institucional(db: Session, nombre_completo: str) -> str:
    partes = unicodedata.normalize('NFKD', nombre_completo.lower())\
        .encode('ascii', 'ignore').decode('ascii')
    partes = re.sub(r'[^a-z\s]', '', partes).split()
    if not partes:
        base = "alumno"
    else:
        nombre = partes[0]
        apellido = partes[1] if len(partes) > 1 else ""
        base = f"{nombre}.{apellido}" if apellido else nombre

    correo = f"{base}@alumno.edupredict.edu"
    sufijo = 1
    while db.query(Estudiante).filter(Estudiante.correo_institucional == correo).first():
        sufijo += 1
        correo = f"{base}{sufijo}@alumno.edupredict.edu"
    return correo


# 1. ESTUDIANTES
@router.post("/estudiantes", response_model=EstudianteOut, dependencies=[Depends(solo_admin)])
def crear_estudiante(estudiante: EstudianteAltaCreate, db: Session = Depends(get_db)):
    grupos_carrera = db.query(Grupo).filter(Grupo.id_carrera == estudiante.id_carrera).all()
    if not grupos_carrera:
        raise HTTPException(
            status_code=400,
            detail="La carrera seleccionada no tiene grupos creados todavía. Crea un grupo primero en Gestión de Grupos."
        )

    conteos = {
        g.id_grupo: db.query(Estudiante).filter(
            Estudiante.id_grupo == g.id_grupo,
            Estudiante.estado == True
        ).count()
        for g in grupos_carrera
    }
    grupo_elegido = min(grupos_carrera, key=lambda g: conteos[g.id_grupo])

    matricula = _generar_matricula(db)
    correo = _generar_correo_institucional(db, estudiante.nombre_completo)

    nuevo_estudiante = Estudiante(
        nombre_completo=estudiante.nombre_completo,
        id_grupo=grupo_elegido.id_grupo,
        datos_socioeconomicos=estudiante.datos_socioeconomicos,
        fecha_ingreso=estudiante.fecha_ingreso,
        contacto_emergencia_nombre=estudiante.contacto_emergencia_nombre,
        contacto_emergencia_telefono=estudiante.contacto_emergencia_telefono,
        edad=estudiante.edad,
        celular=estudiante.celular,
        fotografia_url=estudiante.fotografia_url,
        matricula=matricula,
        correo_institucional=correo
    )
    db.add(nuevo_estudiante)
    db.commit()
    db.refresh(nuevo_estudiante)
    return nuevo_estudiante

@router.get("/estudiantes", response_model=List[EstudianteOut], dependencies=[Depends(todos_los_roles)])
def obtener_estudiantes(
    carrera: Optional[int] = Query(None),
    grupo: Optional[str] = Query(None),
    id_grupo: Optional[int] = Query(None),
    nivel_riesgo: Optional[str] = Query(None),
    skip: int = 0, limit: int = 1000,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol

    query = db.query(Estudiante).filter(Estudiante.estado == True)

    if id_grupo:
        query = query.filter(Estudiante.id_grupo == id_grupo)

    if user_role in ("Tutor", "Mixto") and not (carrera or grupo or id_grupo):
        tutor = db.query(Tutor).filter(Tutor.id_usuario == current_user.id_usuario).first()
        grupos_ids = [g.id_grupo for g in db.query(Grupo).filter(Grupo.id_tutor == (tutor.id_tutor if tutor else -1)).all()]
        query = query.filter(Estudiante.id_grupo.in_(grupos_ids or [-1]))
    elif carrera or grupo:
        query = query.join(Grupo)
        if carrera:
            query = query.filter(Grupo.id_carrera == carrera)
        if grupo:
            query = query.filter(Grupo.nombre_grupo.ilike(f"%{grupo}%"))

    if nivel_riesgo and nivel_riesgo != "":
        query = query.filter(Estudiante.nivel_riesgo == nivel_riesgo)

    estudiantes = query.offset(skip).limit(limit).all()

    ids_grupo = {e.id_grupo for e in estudiantes if e.id_grupo}
    grupos_map = {}
    if ids_grupo:
        grupos_db = db.query(Grupo).filter(Grupo.id_grupo.in_(ids_grupo)).all()
        carreras_ids = {g.id_carrera for g in grupos_db if g.id_carrera}
        carreras_map = {
            c.id_carrera: c.nombre
            for c in db.query(Carrera).filter(Carrera.id_carrera.in_(carreras_ids)).all()
        } if carreras_ids else {}
        grupos_map = {
            g.id_grupo: (g.nombre_grupo, carreras_map.get(g.id_carrera))
            for g in grupos_db
        }

    resultado = []
    for est in estudiantes:
        est_out = EstudianteOut.model_validate(est)
        nombre_grupo, nombre_carrera = grupos_map.get(est.id_grupo, (None, None))
        est_out.nombre_grupo = nombre_grupo
        est_out.nombre_carrera = nombre_carrera
        resultado.append(est_out)

    return resultado

@router.get("/carreras", response_model=List[CarreraOut], dependencies=[Depends(todos_los_roles)])
def obtener_carreras(db: Session = Depends(get_db)):
    return db.query(Carrera).all()

@router.get("/estudiantes/{id}", response_model=EstudianteOut, dependencies=[Depends(todos_los_roles)])
def obtener_estudiante_por_id(id: int, db: Session = Depends(get_db)):
    estudiante = db.query(Estudiante).filter(Estudiante.id_estudiante == id, Estudiante.estado == True).first()
    if not estudiante:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado")
    return estudiante

@router.put("/estudiantes/{id}", response_model=EstudianteOut, dependencies=[Depends(solo_admin)])
def actualizar_estudiante(id: int, est_in: EstudianteCreate, db: Session = Depends(get_db)):
    estudiante = db.query(Estudiante).filter(Estudiante.id_estudiante == id).first()
    if not estudiante:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado")
        
    for key, value in est_in.model_dump().items():
        setattr(estudiante, key, value)
        
    db.commit()
    db.refresh(estudiante)
    return estudiante

@router.delete("/estudiantes/{id}", dependencies=[Depends(solo_admin)])
def eliminar_estudiante_soft(id: int, data: EstudianteBajaIn, db: Session = Depends(get_db)):
    estudiante = db.query(Estudiante).filter(Estudiante.id_estudiante == id).first()
    if not estudiante:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado")
    if not estudiante.estado:
        raise HTTPException(status_code=400, detail="El estudiante ya está dado de baja")

    estudiante.estado = False
    estudiante.motivo_baja = data.motivo_baja
    estudiante.fecha_baja = datetime.utcnow()
    db.commit()
    return {"message": "Estudiante dado de baja correctamente", "motivo_baja": data.motivo_baja}

@router.get("/estudiantes/buscar/{matricula}", response_model=EstudianteOut, dependencies=[Depends(todos_los_roles)])
def buscar_estudiante_por_matricula(matricula: str, db: Session = Depends(get_db)):
    estudiante = db.query(Estudiante).filter(Estudiante.matricula == matricula).first()
    if not estudiante:
        raise HTTPException(status_code=404, detail="No se encontró un estudiante con esa matrícula")
    return estudiante

# 2. GRUPOS
@router.post("/grupos", response_model=GrupoOut, dependencies=[Depends(permitir_gestion_grupos)])
def crear_grupo(
    grupo: GrupoCreate,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol

    if user_role == "Director":
        pertenece = db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario,
            DirectorCarrera.id_carrera == grupo.id_carrera
        ).first()
        if not pertenece:
            raise HTTPException(status_code=403, detail="No puedes crear grupos en una carrera que no diriges.")

    nuevo_grupo = Grupo(**grupo.model_dump())
    db.add(nuevo_grupo)
    db.commit()
    db.refresh(nuevo_grupo)

    carrera = db.query(Carrera).filter(Carrera.id_carrera == nuevo_grupo.id_carrera).first()
    nombre_tutor = None
    if nuevo_grupo.id_tutor:
        tutor_obj = db.query(Tutor).filter(Tutor.id_tutor == nuevo_grupo.id_tutor).first()
        if tutor_obj:
            usuario_tutor = db.query(Usuario).filter(Usuario.id_usuario == tutor_obj.id_usuario).first()
            if usuario_tutor:
                nombre_tutor = usuario_tutor.nombre_completo

    return GrupoOut(
        id_grupo=nuevo_grupo.id_grupo,
        nombre_grupo=nuevo_grupo.nombre_grupo,
        id_carrera=nuevo_grupo.id_carrera,
        nombre_carrera=carrera.nombre if carrera else None,
        nombre_tutor=nombre_tutor,
        cuatrimestre=nuevo_grupo.cuatrimestre,
        id_plan_estudio=nuevo_grupo.id_plan_estudio,
        id_tutor=nuevo_grupo.id_tutor
    )

@router.get("/grupos", response_model=List[GrupoOut], dependencies=[Depends(todos_los_roles)])
def obtener_grupos(skip: int = 0, limit: int = 100, db: Session = Depends(get_db), current_user = Depends(get_current_active_user)):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    query = db.query(Grupo)
    if user_role == "Director":
        mis_carreras = [c.id_carrera for c in db.query(DirectorCarrera).filter(DirectorCarrera.id_usuario == current_user.id_usuario).all()]
        query = query.filter(Grupo.id_carrera.in_(mis_carreras or [-1]))
    grupos = query.offset(skip).limit(limit).all()

    resultado = []
    for g in grupos:
        carrera = db.query(Carrera).filter(Carrera.id_carrera == g.id_carrera).first()
        nombre_tutor = None
        if g.id_tutor:
            tutor_obj = db.query(Tutor).filter(Tutor.id_tutor == g.id_tutor).first()
            if tutor_obj:
                usuario_tutor = db.query(Usuario).filter(Usuario.id_usuario == tutor_obj.id_usuario).first()
                if usuario_tutor:
                    nombre_tutor = usuario_tutor.nombre_completo
        resultado.append(GrupoOut(
            id_grupo=g.id_grupo,
            nombre_grupo=g.nombre_grupo,
            id_carrera=g.id_carrera,
            nombre_carrera=carrera.nombre if carrera else None,
            nombre_tutor=nombre_tutor,
            cuatrimestre=g.cuatrimestre,
            id_plan_estudio=g.id_plan_estudio,
            id_tutor=g.id_tutor
        ))
    return resultado

@router.post("/grupos/{id}/asignar-docente", dependencies=[Depends(solo_admin)])
def asignar_docente_a_grupo(id: int, id_docente: int, db: Session = Depends(get_db)):
    grupo = db.query(Grupo).filter(Grupo.id_grupo == id).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    docente = db.query(Docente).filter(Docente.id_docente == id_docente).first()
    if not docente:
        raise HTTPException(status_code=404, detail="Docente no encontrado")

    usuario = db.query(Usuario).filter(Usuario.id_usuario == docente.id_usuario).first()
    user_role = usuario.rol.value if hasattr(usuario.rol, 'value') else usuario.rol

    tutor = db.query(Tutor).filter(Tutor.id_usuario == docente.id_usuario).first()

    if not tutor:
        usuario.rol = RolEnum.MIXTO
        usuario.token_version = (usuario.token_version or 1) + 1
        nuevo_tutor = Tutor(id_usuario=docente.id_usuario, numero_empleado=f"TUT-{docente.id_usuario}")
        db.add(nuevo_tutor)
        db.flush()
        tutor = nuevo_tutor

    grupo.id_tutor = tutor.id_tutor
    db.commit()
    return {"message": f"{usuario.nombre_completo} asignado como tutor del grupo {id}"}

# 3. MATERIAS
@router.post("/materias", response_model=MateriaOut, dependencies=[Depends(permitir_gestion_materias)])
def crear_materia(materia: MateriaCreate, db: Session = Depends(get_db), current_user = Depends(get_current_active_user)):
    existe = db.query(Materia).filter(Materia.clave_materia == materia.clave_materia).first()
    if existe:
        raise HTTPException(status_code=400, detail="La clave de materia ya está en uso")

    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role == "Director":
        mis_carreras = [c.id_carrera for c in db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario
        ).all()]
        if materia.id_carrera not in mis_carreras:
            raise HTTPException(status_code=403, detail="No puedes crear materias para una carrera que no diriges.")

    nueva_materia = Materia(**materia.model_dump())
    db.add(nueva_materia)
    db.commit()
    db.refresh(nueva_materia)
    return _armar_materia_salida(nueva_materia, db)

@router.put("/materias/{id}", response_model=MateriaOut, dependencies=[Depends(permitir_gestion_materias)])
def actualizar_materia(id: int, data: MateriaUpdate, db: Session = Depends(get_db), current_user = Depends(get_current_active_user)):
    materia = db.query(Materia).filter(Materia.id_materia == id).first()
    if not materia:
        raise HTTPException(status_code=404, detail="Materia no encontrada")

    if data.clave_materia and data.clave_materia != materia.clave_materia:
        existe = db.query(Materia).filter(Materia.clave_materia == data.clave_materia).first()
        if existe:
            raise HTTPException(status_code=400, detail="La clave de materia ya está en uso")

    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role == "Director":
        mis_carreras = [c.id_carrera for c in db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario
        ).all()]
        id_carrera_objetivo = data.id_carrera if data.id_carrera is not None else materia.id_carrera
        if id_carrera_objetivo not in mis_carreras:
            raise HTTPException(status_code=403, detail="No puedes modificar materias fuera de tu carrera.")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(materia, key, value)

    db.commit()
    db.refresh(materia)
    return _armar_materia_salida(materia, db)

@router.get("/materias", response_model=List[MateriaOut], dependencies=[Depends(todos_los_roles)])
def obtener_materias(
    carrera: Optional[int] = Query(None),
    skip: int = 0, limit: int = 100,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    query = db.query(Materia).filter(Materia.estado == True)

    if user_role == "Director":
        mis_carreras = [c.id_carrera for c in db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario
        ).all()]
        query = query.filter(Materia.id_carrera.in_(mis_carreras or [-1]))
    elif carrera:
        query = query.filter(Materia.id_carrera == carrera)

    materias = query.offset(skip).limit(limit).all()
    return [_armar_materia_salida(m, db) for m in materias]

@router.get("/grupos/{id_grupo}/mis-materias", response_model=List[MateriaOut])
def materias_del_docente_en_grupo(
    id_grupo: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol

    if user_role == "Administrador":
        materia_ids = db.query(Horario.id_materia).filter(Horario.id_grupo == id_grupo).distinct()
    else:
        docente = db.query(Docente).filter(Docente.id_usuario == current_user.id_usuario).first()
        if not docente:
            return []
        materia_ids = db.query(Horario.id_materia).filter(
            Horario.id_grupo == id_grupo,
            Horario.id_docente == docente.id_docente
        ).distinct()

    ids = [m[0] for m in materia_ids]
    return db.query(Materia).filter(Materia.id_materia.in_(ids)).all()

@router.get("/mis-clases", response_model=List[ClaseDocenteOut])
def obtener_mis_clases(
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role != "Docente":
        raise HTTPException(status_code=403, detail="Solo disponible para el rol Docente")

    docente = db.query(Docente).filter(Docente.id_usuario == current_user.id_usuario).first()
    if not docente:
        raise HTTPException(status_code=400, detail="Tu cuenta no tiene un perfil de docente asociado")

    horarios = db.query(Horario).filter(Horario.id_docente == docente.id_docente).all()

    resultado = []
    for h in horarios:
        grupo = db.query(Grupo).filter(Grupo.id_grupo == h.id_grupo).first()
        materia = db.query(Materia).filter(Materia.id_materia == h.id_materia).first()
        if not grupo or not materia:
            continue

        num_alumnos = db.query(Estudiante).filter(
            Estudiante.id_grupo == grupo.id_grupo,
            Estudiante.estado == True
        ).count()

        resultado.append(ClaseDocenteOut(
            id_horario=h.id_horario,
            id_grupo=grupo.id_grupo,
            nombre_grupo=grupo.nombre_grupo,
            id_materia=materia.id_materia,
            nombre_materia=materia.nombre_materia,
            dia_semana=h.dia_semana.value if h.dia_semana else "",
            hora_inicio=h.hora_inicio.strftime("%H:%M") if h.hora_inicio else "",
            hora_fin=h.hora_fin.strftime("%H:%M") if h.hora_fin else "",
            num_alumnos=num_alumnos
        ))

    orden_dias = {"LUNES": 0, "MARTES": 1, "MIERCOLES": 2, "JUEVES": 3, "VIERNES": 4}
    resultado.sort(key=lambda c: (orden_dias.get(c.dia_semana.upper().replace("É", "E"), 9), c.hora_inicio))

    return resultado

# 4. PERIODOS ACADÉMICOS
@router.post("/periodos", response_model=PeriodoOut, dependencies=[Depends(solo_admin)])
def crear_periodo(periodo: PeriodoCreate, db: Session = Depends(get_db)):
    nuevo_periodo = Periodo(**periodo.model_dump())
    db.add(nuevo_periodo)
    db.commit()
    db.refresh(nuevo_periodo)
    return nuevo_periodo

@router.get("/periodos", response_model=List[PeriodoOut], dependencies=[Depends(todos_los_roles)])
def obtener_periodos(db: Session = Depends(get_db)):
    return db.query(Periodo).all()

@router.get("/grupos/plantilla-importacion", dependencies=[Depends(permitir_gestion_grupos)])
def descargar_plantilla_importacion():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    ws.append(["Matricula"])
    ws.append(["MAT-2024001"])
    ws.append(["MAT-2024002"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_importacion_alumnos.xlsx"}
    )


@router.post("/grupos/{id_grupo}/importar-alumnos", dependencies=[Depends(permitir_gestion_grupos)])
def importar_alumnos_a_grupo(
    id_grupo: int,
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    grupo = db.query(Grupo).filter(Grupo.id_grupo == id_grupo).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role == "Director":
        pertenece = db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario,
            DirectorCarrera.id_carrera == grupo.id_carrera
        ).first()
        if not pertenece:
            raise HTTPException(status_code=403, detail="No puedes importar alumnos a un grupo de una carrera que no diriges.")

    if not archivo.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx o .xls")

    try:
        contenido = archivo.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. Verifica el formato.")

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    if not filas:
        raise HTTPException(status_code=400, detail="El archivo no contiene registros (fila 2 en adelante).")

    actualizados = []
    no_encontrados = []
    ya_asignados_otro_grupo = []

    for fila in filas:
        if not fila or not fila[0]:
            continue
        matricula = str(fila[0]).strip()
        if not matricula:
            continue

        estudiante = db.query(Estudiante).filter(Estudiante.matricula == matricula).first()
        if not estudiante:
            no_encontrados.append(matricula)
            continue

        if estudiante.id_grupo and estudiante.id_grupo != id_grupo:
            ya_asignados_otro_grupo.append(matricula)

        estudiante.id_grupo = id_grupo
        actualizados.append(matricula)

    db.commit()

    return {
        "message": f"{len(actualizados)} alumno(s) vinculados al grupo '{grupo.nombre_grupo}'.",
        "actualizados": actualizados,
        "no_encontrados": no_encontrados,
        "reasignados_desde_otro_grupo": ya_asignados_otro_grupo
    }

def _hay_colision_horario(db: Session, id_docente: int, id_aula, dia_semana, hora_inicio, hora_fin, excluir_id_horario=None):
    query = db.query(Horario).filter(Horario.dia_semana == dia_semana)
    if excluir_id_horario:
        query = query.filter(Horario.id_horario != excluir_id_horario)

    candidatos = query.filter(
        (Horario.id_docente == id_docente) | (Horario.id_aula == id_aula if id_aula else False)
    ).all()

    for h in candidatos:
        if h.hora_inicio is None or h.hora_fin is None:
            continue
        traslape = h.hora_inicio < hora_fin and hora_inicio < h.hora_fin
        if not traslape:
            continue
        if h.id_docente == id_docente:
            return f"El docente ya tiene clase de {h.hora_inicio.strftime('%H:%M')} a {h.hora_fin.strftime('%H:%M')} el {dia_semana.value}."
        if id_aula and h.id_aula == id_aula:
            return f"El aula ya está ocupada de {h.hora_inicio.strftime('%H:%M')} a {h.hora_fin.strftime('%H:%M')} el {dia_semana.value}."
    return None


@router.get("/docentes-catalogo", response_model=List[DocenteCatalogoOut], dependencies=[Depends(todos_los_roles)])
def catalogo_docentes(db: Session = Depends(get_db)):
    docentes = db.query(Docente, Usuario)\
        .join(Usuario, Docente.id_usuario == Usuario.id_usuario)\
        .filter(Usuario.estado == True).all()
    resultado = []
    for d, u in docentes:
        nombre_parts = u.nombre_completo.split()
        nombre = nombre_parts[0] if nombre_parts else ""
        apellidos = " ".join(nombre_parts[1:]) if len(nombre_parts) > 1 else ""
        resultado.append(DocenteCatalogoOut(
            id=d.id_docente,
            nombre=nombre,
            apellidos=apellidos,
            correo=u.correo_institucional or ""
        ))
    return resultado


@router.get("/grupos/{id_grupo}/horarios", response_model=List[HorarioOut], dependencies=[Depends(todos_los_roles)])
def obtener_horarios_grupo(id_grupo: int, db: Session = Depends(get_db)):
    horarios = db.query(Horario).filter(Horario.id_grupo == id_grupo).all()
    resultado = []
    for h in horarios:
        materia = db.query(Materia).filter(Materia.id_materia == h.id_materia).first()
        docente = db.query(Docente).filter(Docente.id_docente == h.id_docente).first()
        nombre_docente = None
        if docente:
            usuario = db.query(Usuario).filter(Usuario.id_usuario == docente.id_usuario).first()
            nombre_docente = usuario.nombre_completo if usuario else None

        resultado.append(HorarioOut(
            id_horario=h.id_horario,
            id_grupo=h.id_grupo,
            id_materia=h.id_materia,
            nombre_materia=materia.nombre_materia if materia else None,
            id_docente=h.id_docente,
            nombre_docente=nombre_docente,
            id_aula=h.id_aula,
            dia_semana=h.dia_semana.value if h.dia_semana else "",
            hora_inicio=h.hora_inicio.strftime("%H:%M") if h.hora_inicio else "",
            hora_fin=h.hora_fin.strftime("%H:%M") if h.hora_fin else ""
        ))
    return resultado


@router.post("/horarios/lote", dependencies=[Depends(permitir_gestion_grupos)])
def crear_horarios_lote(
    data: HorarioBulkCreate,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    grupo = db.query(Grupo).filter(Grupo.id_grupo == data.id_grupo).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role == "Director":
        pertenece = db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario,
            DirectorCarrera.id_carrera == grupo.id_carrera
        ).first()
        if not pertenece:
            raise HTTPException(status_code=403, detail="No puedes editar horarios de un grupo de una carrera que no diriges.")

    creados = []
    errores = []
    lote_verificado = []  # para detectar colisiones dentro del mismo lote enviado

    for idx, h in enumerate(data.horarios):
        dia_enum = getattr(DiaSemanaEnum, h.dia_semana.upper().replace('É', 'E'), None)
        if not dia_enum:
            errores.append({"fila": idx + 1, "error": f"Día inválido: {h.dia_semana}"})
            continue
        if h.hora_fin <= h.hora_inicio:
            errores.append({"fila": idx + 1, "error": "La hora de fin debe ser posterior a la hora de inicio."})
            continue

        conflicto = _hay_colision_horario(db, h.id_docente, h.id_aula, dia_enum, h.hora_inicio, h.hora_fin)
        if conflicto:
            errores.append({"fila": idx + 1, "error": conflicto})
            continue

        conflicto_lote = None
        for b in lote_verificado:
            if b["dia"] == dia_enum and b["hora_inicio"] < h.hora_fin and h.hora_inicio < b["hora_fin"]:
                if b["id_docente"] == h.id_docente:
                    conflicto_lote = "Choca con otra fila de esta misma solicitud (mismo docente, mismo horario)."
                elif h.id_aula and b["id_aula"] == h.id_aula:
                    conflicto_lote = "Choca con otra fila de esta misma solicitud (misma aula, mismo horario)."
        if conflicto_lote:
            errores.append({"fila": idx + 1, "error": conflicto_lote})
            continue

        nuevo = Horario(
            id_grupo=data.id_grupo,
            id_materia=h.id_materia,
            id_docente=h.id_docente,
            id_aula=h.id_aula,
            dia_semana=dia_enum,
            hora_inicio=h.hora_inicio,
            hora_fin=h.hora_fin
        )
        db.add(nuevo)
        db.flush()
        creados.append(nuevo.id_horario)
        lote_verificado.append({
            "dia": dia_enum, "hora_inicio": h.hora_inicio, "hora_fin": h.hora_fin,
            "id_docente": h.id_docente, "id_aula": h.id_aula
        })

    if errores and not creados:
        db.rollback()
        raise HTTPException(status_code=400, detail={"message": "No se pudo guardar ningún horario.", "errores": errores})

    db.commit()
    return {"message": f"{len(creados)} horario(s) creados correctamente.", "creados": creados, "errores": errores}


@router.delete("/horarios/{id_horario}", dependencies=[Depends(permitir_gestion_grupos)])
def eliminar_horario(
    id_horario: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    h = db.query(Horario).filter(Horario.id_horario == id_horario).first()
    if not h:
        raise HTTPException(status_code=404, detail="Horario no encontrado")

    user_role = current_user.rol.value if hasattr(current_user.rol, 'value') else current_user.rol
    if user_role == "Director":
        grupo = db.query(Grupo).filter(Grupo.id_grupo == h.id_grupo).first()
        pertenece = db.query(DirectorCarrera).filter(
            DirectorCarrera.id_usuario == current_user.id_usuario,
            DirectorCarrera.id_carrera == grupo.id_carrera if grupo else None
        ).first()
        if not pertenece:
            raise HTTPException(status_code=403, detail="No puedes eliminar horarios de un grupo de una carrera que no diriges.")

    db.delete(h)
    db.commit()
    return {"message": "Horario eliminado correctamente."}

permitir_historial = RoleChecker(["Administrador", "Psicopedagogia", "Tutor", "Director", "Mixto"])

@router.get("/estudiantes/{id_estudiante}/historial-previo", response_model=List[HistorialAcademicoPrevioOut], dependencies=[Depends(permitir_historial)])
def obtener_historial_previo(id_estudiante: int, db: Session = Depends(get_db)):
    return db.query(HistorialAcademicoPrevio)\
        .filter(HistorialAcademicoPrevio.id_estudiante == id_estudiante)\
        .order_by(HistorialAcademicoPrevio.fecha_registro.desc())\
        .all()

@router.post("/estudiantes/{id_estudiante}/historial-previo", response_model=HistorialAcademicoPrevioOut, dependencies=[Depends(solo_admin)])
def crear_historial_previo(id_estudiante: int, data: HistorialAcademicoPrevioCreate, db: Session = Depends(get_db)):
    estudiante = db.query(Estudiante).filter(Estudiante.id_estudiante == id_estudiante).first()
    if not estudiante:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado")

    nuevo = HistorialAcademicoPrevio(id_estudiante=id_estudiante, **data.model_dump())
    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)
    return nuevo

@router.delete("/historial-previo/{id_historial}", dependencies=[Depends(solo_admin)])
def eliminar_historial_previo(id_historial: int, db: Session = Depends(get_db)):
    registro = db.query(HistorialAcademicoPrevio).filter(HistorialAcademicoPrevio.id_historial == id_historial).first()
    if not registro:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    db.delete(registro)
    db.commit()
    return {"message": "Registro eliminado correctamente"}